"""HR веб-панель для просмотра и управления кандидатами."""

from __future__ import annotations

import asyncio
import io
import logging
import os
import sys
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook

from config import (
    AI_BASE_URL,
    AI_ENABLED,
    BOT_TOKEN,
    HR_PANEL_HOST,
    HR_PANEL_PASSWORD,
    HR_PANEL_PORT,
    HR_PANEL_SECRET_KEY,
    POSITIONS,
    STATUS_OPTIONS,
)
from database import backup_db, get_all_candidates, get_candidate_by_id, init_db, update_hr_status

app = Flask(
    __name__,
    template_folder=str(Path(__file__).resolve().parent / "templates"),
    static_folder=str(Path(__file__).resolve().parent / "static"),
)
app.secret_key = HR_PANEL_SECRET_KEY

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

def login_required(f):
    """Декоратор: требует авторизации для доступа к странице."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


@app.route("/login", methods=["GET", "POST"])
def login():
    """Страница входа в HR-панель."""
    if request.method == "POST":
        password = request.form.get("password", "")
        if password == HR_PANEL_PASSWORD:
            session["logged_in"] = True
            flash("Вы вошли в систему", "success")
            return redirect(url_for("candidates_list"))
        flash("Неверный пароль", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    """Выход из HR-панели."""
    session.clear()
    flash("Вы вышли из системы", "info")
    return redirect(url_for("login"))


def check_ai_status() -> bool:
    """Проверить доступность AI-сервера (Ollama) быстрым запросом."""
    import urllib.request
    try:
        req = urllib.request.Request(f"{AI_BASE_URL}/models")
        with urllib.request.urlopen(req, timeout=1.5):
            return True
    except Exception:
        return False


# ---------------------------------------------------------------------------
# Candidates list
# ---------------------------------------------------------------------------

@app.route("/")
@login_required
def candidates_list():
    """Главная страница — таблица кандидатов с фильтрами."""
    sort_by = request.args.get("sort_by", "created_at")
    sort_order = request.args.get("sort_order", "DESC")
    status_filter = request.args.get("status") or None
    hr_status_filter = request.args.get("hr_status") or None
    position_filter = request.args.get("position") or None
    min_experience = request.args.get("min_experience", type=float)
    lang_filter = request.args.get("lang") or None
    search_name = request.args.get("search") or None

    candidates = get_all_candidates(
        sort_by=sort_by,
        sort_order=sort_order,
        status_filter=status_filter,
        hr_status_filter=hr_status_filter,
        position_filter=position_filter,
        min_experience=min_experience,
        lang_filter=lang_filter,
        search_name=search_name,
    )

    next_order = "ASC" if sort_order == "DESC" else "DESC"

    return render_template(
        "candidates.html",
        candidates=candidates,
        status_options=STATUS_OPTIONS,
        positions=POSITIONS,
        ai_enabled=AI_ENABLED,
        ai_online=check_ai_status() if AI_ENABLED else False,
        sort_by=sort_by,
        sort_order=sort_order,
        next_order=next_order,
        current_status=status_filter or "",
        current_hr_status=hr_status_filter or "",
        current_position=position_filter or "",
        current_experience=min_experience if min_experience is not None else "",
        current_lang=lang_filter or "",
        current_search=search_name or "",
        total=len(candidates),
    )


# ---------------------------------------------------------------------------
# Invite to interview
# ---------------------------------------------------------------------------

def send_telegram_message_sync(chat_id: int, text: str) -> bool:
    """Отправить сообщение кандидату через Telegram Bot API."""
    async def _send():
        from aiogram import Bot
        bot = Bot(token=BOT_TOKEN)
        try:
            await bot.send_message(chat_id=chat_id, text=text)
            return True
        except Exception:
            logger.exception("Ошибка отправки сообщения tg_id=%s", chat_id)
            return False
        finally:
            await bot.session.close()

    try:
        loop = asyncio.get_event_loop()
        if loop.is_running():
            import concurrent.futures
            with concurrent.futures.ThreadPoolExecutor() as pool:
                future = pool.submit(asyncio.run, _send())
                return future.result(timeout=15)
        return loop.run_until_complete(_send())
    except RuntimeError:
        return asyncio.run(_send())


@app.route("/invite/<int:candidate_id>", methods=["GET", "POST"])
@login_required
def invite(candidate_id: int):
    """Страница приглашения кандидата на собеседование."""
    candidate = get_candidate_by_id(candidate_id)
    if not candidate:
        flash("Кандидат не найден", "danger")
        return redirect(url_for("candidates_list"))

    if request.method == "POST":
        message_text = request.form.get("message", "").strip()
        if not message_text:
            message_text = (
                f"Здравствуйте, {candidate['full_name']}!\n\n"
                "Мы рассмотрели вашу анкету и хотели бы пригласить вас "
                "на собеседование.\n\n"
                "Наш HR-менеджер свяжется с вами для уточнения деталей."
            )

        success = send_telegram_message_sync(candidate["telegram_id"], message_text)
        update_hr_status(candidate_id, "приглашён")

        if success:
            flash(f"Приглашение отправлено: {candidate['full_name']}", "success")
        else:
            flash("Статус обновлён, но сообщение не доставлено (проверьте токен бота)", "warning")

        return redirect(url_for("candidates_list"))

    default_message = (
        f"Здравствуйте, {candidate['full_name']}!\n\n"
        "Мы рассмотрели вашу анкету и хотели бы пригласить вас "
        "на собеседование.\n\n"
        "Наш HR-менеджер свяжется с вами для уточнения деталей."
    )

    return render_template(
        "invite.html",
        candidate=candidate,
        default_message=default_message,
    )


# ---------------------------------------------------------------------------
# Reject candidate
# ---------------------------------------------------------------------------

@app.route("/reject/<int:candidate_id>", methods=["GET", "POST"])
@login_required
def reject(candidate_id: int):
    """Страница отказа кандидату."""
    candidate = get_candidate_by_id(candidate_id)
    if not candidate:
        flash("Кандидат не найден", "danger")
        return redirect(url_for("candidates_list"))

    if request.method == "POST":
        message_text = request.form.get("message", "").strip()
        if not message_text:
            message_text = (
                f"Здравствуйте, {candidate['full_name']}!\n\n"
                "Благодарим вас за интерес к нашей компании и время, "
                "потраченное на заполнение анкеты.\n\n"
                "К сожалению, на данный момент мы не готовы сделать вам предложение. "
                "Мы сохраним ваше резюме и свяжемся, если появится подходящая позиция.\n\n"
                "Желаем удачи!"
            )

        success = send_telegram_message_sync(candidate["telegram_id"], message_text)
        update_hr_status(candidate_id, "отказано")

        if success:
            flash(f"Отказ отправлен: {candidate['full_name']}", "warning")
        else:
            flash("Статус обновлён, но сообщение не доставлено (проверьте токен бота)", "warning")

        return redirect(url_for("candidates_list"))

    default_message = (
        f"Здравствуйте, {candidate['full_name']}!\n\n"
        "Благодарим вас за интерес к нашей компании и время, "
        "потраченное на заполнение анкеты.\n\n"
        "К сожалению, на данный момент мы не готовы сделать вам предложение. "
        "Мы сохраним ваше резюме и свяжемся, если появится подходящая позиция.\n\n"
        "Желаем удачи!"
    )

    return render_template(
        "reject.html",
        candidate=candidate,
        default_message=default_message,
    )


# ---------------------------------------------------------------------------
# Export to Excel
# ---------------------------------------------------------------------------

@app.route("/export")
@login_required
def export_excel():
    """Экспорт кандидатов в Excel с учётом текущих фильтров."""
    status_filter = request.args.get("status") or None
    hr_status_filter = request.args.get("hr_status") or None
    position_filter = request.args.get("position") or None
    min_experience = request.args.get("min_experience", type=float)
    lang_filter = request.args.get("lang") or None
    search_name = request.args.get("search") or None

    candidates = get_all_candidates(
        status_filter=status_filter,
        hr_status_filter=hr_status_filter,
        position_filter=position_filter,
        min_experience=min_experience,
        lang_filter=lang_filter,
        search_name=search_name,
    )

    position_titles = {p["id"]: p["title"] for p in POSITIONS}

    wb = Workbook()
    ws = wb.active
    ws.title = "Кандидаты"

    headers = [
        "ID", "Telegram ID", "Username", "ФИО",
        "Контакт", "Вакансия", "Опыт (лет)", "Навыки",
        "Статус", "HR-статус", "AI-оценка", "AI-саммари",
        "Готов приступить", "Дата заявки",
    ]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).font = ws.cell(row=1, column=col).font.copy(bold=True)

    for c in candidates:
        ws.append([
            c["id"],
            c["telegram_id"],
            c.get("telegram_username", ""),
            c["full_name"],
            c["email_or_phone"],
            position_titles.get(c.get("position"), c.get("position", "")),
            c["experience_years"],
            c["programming_languages"],
            c["status"],
            c.get("hr_status") or "на рассмотрении",
            c.get("ai_score") or "",
            c.get("ai_summary") or "",
            c["start_date"],
            c["created_at"],
        ])

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value) if cell.value else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"candidates_{timestamp}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# Backup
# ---------------------------------------------------------------------------

@app.route("/backup", methods=["POST"])
@login_required
def do_backup():
    """Создать резервную копию базы данных."""
    try:
        path = backup_db()
        flash(f"Бэкап создан: {path.name}", "success")
    except Exception:
        logger.exception("Ошибка создания бэкапа")
        flash("Ошибка создания бэкапа", "danger")
    return redirect(url_for("candidates_list"))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )
    init_db()
    logger.info("HR-панель запущена: http://%s:%s", HR_PANEL_HOST, HR_PANEL_PORT)
    debug_mode = os.getenv("FLASK_DEBUG", "0") == "1"
    app.run(host=HR_PANEL_HOST, port=HR_PANEL_PORT, debug=debug_mode)
